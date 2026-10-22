# -*- coding: utf-8 -*-
"""按分项编号生成 JA-001 / JA-002 / JA-003 资料参数表。

测站匹配保持 R1：互通走廊只认全称并锁池，邻近上限 2000m。
不启用 3km 兜底。缺导线或水准的数据行整行标黄。
"""
from __future__ import annotations

import hashlib
import re
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parents[2]
CANONICAL_TEMPLATE = REPO_ROOT / "share" / "参数样板.xlsx"
MAX_NEAR_M = 2000
YELLOW = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
PROJECT_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9_-]{0,31}")
DEF_RE = re.compile(
    r"D\s*[，,]\s*E\s*[，,]\s*F列分别写上"
    r"[“\"]([^”\"]+)[”\"]\s*[，,]\s*"
    r"[“\"]([^”\"]+)[”\"]\s*[，,]\s*"
    r"[“\"]([^”\"]+)[”\"]"
)
CONST_RE = re.compile(
    r"AD至AF列分别写上\s*"
    r"([0-9]+(?:\.[0-9]+)?)\s*[，,]\s*"
    r"([0-9]+(?:\.[0-9]+)?)\s*[，,]\s*"
    r"([0-9]+(?:\.[0-9]+)?)"
)

STAKE_RE = re.compile(r"(?<![A-Za-z])([A-Za-z]{0,3}K)(\d+)\+(\d+(?:\.\d+)?)", re.IGNORECASE)

# R1 互通走廊名（只用规范全称，避免「白岭大桥/分离」等短名误伤）
CORRIDOR_CANON = [
    "白岭互通",
    "石坳互通",
    "港口互通",
    "修水西枢纽",
    "铜鼓枢纽",
]

SIDE_PHRASES = [
    ("路线右侧", "右幅"),
    ("路线左侧", "左幅"),
    ("匝道右侧", "右幅"),
    ("匝道左侧", "左幅"),
    ("主线右侧", "右幅"),
    ("主线左侧", "左幅"),
    ("路基右侧", "右幅"),
    ("路基左侧", "左幅"),
    ("行车方向右侧", "右幅"),
    ("行车方向左侧", "左幅"),
]

SIDE_STRIP_RES = [
    re.compile(r"(左线|右线)?行车方向右侧$"),
    re.compile(r"(左线|右线)?行车方向左侧$"),
    re.compile(r"路线右侧$"),
    re.compile(r"路线左侧$"),
    re.compile(r"匝道右侧$"),
    re.compile(r"匝道左侧$"),
    re.compile(r"主线右侧$"),
    re.compile(r"主线左侧$"),
    re.compile(r"路基右侧$"),
    re.compile(r"路基左侧$"),
    re.compile(r"改路两侧$"),
    re.compile(r"改路右侧$"),
    re.compile(r"改路左侧$"),
    re.compile(r"两侧$"),
    re.compile(r"右侧$"),
    re.compile(r"左侧$"),
]


def stake_to_int(km: str, m: str) -> int:
    return int(round(float(km) * 1000.0 + float(m)))


def round10(v: int) -> int:
    return int(round(v / 10.0) * 10)


def parse_stakes(text: str):
    out = []
    for m in STAKE_RE.finditer(text or ""):
        prefix = m.group(1).upper()
        num = stake_to_int(m.group(2), m.group(3))
        out.append((prefix, num, m.group(0)))
    return out


def clean_b(raw: str, eng_prefix: str) -> str:
    s = str(raw).strip()
    if s.startswith(eng_prefix):
        s = s[len(eng_prefix) :]
    changed = True
    while changed:
        changed = False
        for rx in SIDE_STRIP_RES:
            ns = rx.sub("", s)
            if ns != s:
                s = ns
                changed = True
    return s.strip()


def detect_side(raw: str):
    s = str(raw)
    best = None
    best_pos = -1
    for phrase, label in SIDE_PHRASES:
        pos = s.rfind(phrase)
        if pos > best_pos:
            best_pos = pos
            best = label
    if best:
        return best, ""

    rpos = s.rfind("右侧")
    lpos = s.rfind("左侧")
    if rpos >= 0 or lpos >= 0:
        if rpos > lpos:
            return "右幅", ""
        if lpos > rpos:
            return "左幅", ""

    if "中央分隔带" in s or "中分带" in s or ("中央" in s and "左" not in s and "右" not in s):
        return "左幅", "【推断】中央/中分带无左右侧字样，取左幅"

    if "右" in s and "左" not in s:
        return "右幅", "【推断】仅见「右」"
    if "左" in s and "右" not in s:
        return "左幅", "【推断】仅见「左」"
    if "左" in s and "右" in s:
        last_l = s.rfind("左")
        last_r = s.rfind("右")
        if last_r > last_l:
            return "右幅", "【推断】左右皆有，取靠后「右」"
        return "左幅", "【推断】左右皆有，取靠后「左」"

    return "", "【无法判定】无左右侧信息"


def pick_ab_ja001(span: int) -> int:
    if span <= 50:
        return 10
    if span <= 70:
        return 20
    return 40


def pick_ab_ja002(span: int) -> int:
    if span <= 50:
        return 10
    if span <= 70:
        return 20
    if span <= 200:
        return 40
    return 100


def compute_zab(nums: list[int], pick_ab):
    if not nums:
        return None, None, None, False, "无桩号"
    if len(nums) == 1:
        z = round10(nums[0])
        adjusted = z != nums[0]
        note = f"单桩号取整10: {nums[0]}->{z}" if adjusted else ""
        return z, z, 0, adjusted, note

    z_raw, aa_raw = nums[0], nums[1]
    z, aa = round10(z_raw), round10(aa_raw)
    adjusted = (z != z_raw) or (aa != aa_raw)
    notes = []
    if adjusted:
        notes.append(f"起止取整10: {z_raw}/{aa_raw}->{z}/{aa}")

    if z == aa:
        return z, aa, 0, adjusted, "; ".join(notes) or ""

    ascending = aa >= z
    span = abs(aa - z)
    ab = pick_ab(span)
    if span % ab != 0:
        new_span = int(round(span / float(ab)) * ab)
        if new_span == 0:
            new_span = ab
        aa = z + new_span if ascending else z - new_span
        span2 = abs(aa - z)
        ab2 = pick_ab(span2)
        if span2 % ab2 != 0:
            new_span = (span2 // ab2) * ab2
            if new_span == 0:
                new_span = ab2
            aa = z + new_span if ascending else z - new_span
            ab = ab2
        else:
            ab = ab2
        adjusted = True
        notes.append(f"微调使可被AB整除: AB={ab}, AA={aa}")
    return z, aa, ab, adjusted, "; ".join(notes)


def load_points(ws, need_xy: bool):
    pts = []
    for row in ws.iter_rows(min_row=3, max_col=5, values_only=True):
        name, x, y, z, pos = row[0], row[1], row[2], row[3], row[4]
        if not name:
            continue
        stakes = parse_stakes(str(pos or ""))
        if not stakes:
            continue
        prefix, num, _ = stakes[0]
        if need_xy and (x is None or y is None):
            continue
        if (not need_xy) and z is None:
            continue
        pts.append(
            {
                "name": str(name).strip(),
                "x": float(x) if x is not None else None,
                "y": float(y) if y is not None else None,
                "z": float(z) if z is not None else None,
                "prefix": prefix,
                "num": num,
                "pos": str(pos or ""),
            }
        )
    return pts


def _within(points, target_num, max_m=MAX_NEAR_M):
    return [p for p in points if abs(p["num"] - target_num) <= max_m]


def extract_corridor(text: str):
    """从台账 B / 测站位置描述提取规范互通名；仅匹配全称。"""
    s = str(text or "")
    for name in CORRIDOR_CANON:
        if name in s:
            return name
    return None


def point_corridor(pt) -> str | None:
    return extract_corridor(pt.get("pos", ""))


def corridor_pool(points, corridor: str):
    """位置描述含该互通名（含「辐射整个××」）的测站。"""
    return [p for p in points if corridor in (p.get("pos") or "")]


def corridor_k_median(points, corridor: str):
    """走廊内主线 K 桩号中位数；无 K 则用池内全部桩号中位数。"""
    pool = corridor_pool(points, corridor)
    k_nums = [p["num"] for p in pool if p["prefix"] == "K"]
    nums = k_nums if k_nums else [p["num"] for p in pool]
    if not nums:
        return None
    nums = sorted(nums)
    return float(nums[len(nums) // 2])


def nearest_two(points, target_num, prefer_prefix, max_m=MAX_NEAR_M):
    if not points:
        return None, None
    same = [p for p in points if p["prefix"] == prefer_prefix] if prefer_prefix else []
    pool = same if len(same) >= 2 else points
    ranked = sorted(_within(pool, target_num, max_m), key=lambda p: (abs(p["num"] - target_num), p["num"]))
    if len(ranked) < 2:
        ranked = sorted(_within(points, target_num, max_m), key=lambda p: (abs(p["num"] - target_num), p["num"]))
    if len(ranked) < 2:
        return None, None
    return ranked[0], ranked[1]


def nearest_one(points, target_num, prefer_prefix, max_m=MAX_NEAR_M):
    if not points:
        return None
    same = [p for p in points if p["prefix"] == prefer_prefix] if prefer_prefix else []
    pool = same if same else points
    cand = _within(pool, target_num, max_m)
    if not cand:
        cand = _within(points, target_num, max_m)
    if not cand:
        return None
    return min(cand, key=lambda p: (abs(p["num"] - target_num), p["num"]))


def nearest_two_any(points, target_num):
    """无 2km 限制：按距 target 最近取 2 个（用于无桩号+互通走廊）。"""
    if len(points) < 2:
        return None, None
    ranked = sorted(points, key=lambda p: (abs(p["num"] - target_num), p["num"]))
    return ranked[0], ranked[1]


def nearest_one_any(points, target_num):
    if not points:
        return None
    return min(points, key=lambda p: (abs(p["num"] - target_num), p["num"]))


def match_stations(raw_b, prefix, target, traverse, levels, corridor_medians):
    """R1+R2 测站匹配。

    返回 (station, back, level, meta) 其中 meta 含:
      corridor, path (corridor_stake|corridor_no_stake|global), miss_t, miss_l
    """
    corridor = extract_corridor(raw_b)
    meta = {
        "corridor": corridor,
        "path": "global",
        "miss_t": False,
        "miss_l": False,
        "no_stake_corridor_fill": False,
        "target_used": target,
    }

    if corridor:
        t_pool = corridor_pool(traverse, corridor)
        l_pool = corridor_pool(levels, corridor)
        if target is not None:
            # 有桩号：池内同前缀优先，再池内任意前缀，仍 ≤2km
            station, back = nearest_two(t_pool, target, prefix)
            level = nearest_one(l_pool, target, prefix)
            meta["path"] = "corridor_stake"
            # 池内按台账桩号不足时（典型：匝道 AK/CK 而池多为「辐射」主线 K）：
            # R1「池内不足→走廊主线K中位数」——在池内按中位数取点，距离相对中位数仍≤2km
            if (station is None or back is None or level is None) and len(t_pool) >= 2 and len(l_pool) >= 1:
                med = corridor_medians.get(corridor)
                if med is None:
                    med = corridor_k_median(traverse, corridor)
                if med is not None:
                    s2, b2 = nearest_two(t_pool, med, "K")
                    if s2 is None or b2 is None:
                        s2, b2 = nearest_two(t_pool, med, None)
                    lv2 = nearest_one(l_pool, med, "K")
                    if lv2 is None:
                        lv2 = nearest_one(l_pool, med, None)
                    if s2 and b2 and lv2:
                        station, back, level = s2, b2, lv2
                        meta["path"] = "corridor_stake_median"
                        meta["target_used"] = med
            if station is None or back is None:
                meta["miss_t"] = True
            if level is None:
                meta["miss_l"] = True
            return station, back, level, meta

        # 无桩号 + 有互通名：用走廊主线 K 中位数作目标，池内取最近（不套 2km）
        med = corridor_medians.get(corridor)
        if med is None:
            med = corridor_k_median(traverse, corridor)
        if med is None:
            meta["miss_t"] = meta["miss_l"] = True
            meta["path"] = "corridor_no_stake"
            return None, None, None, meta

        meta["target_used"] = med
        meta["path"] = "corridor_no_stake"
        if len(t_pool) >= 2 and len(l_pool) >= 1:
            station, back = nearest_two_any(t_pool, med)
            level = nearest_one_any(l_pool, med)
            meta["no_stake_corridor_fill"] = True
            return station, back, level, meta

        # 池太小：仍黄
        station, back = (nearest_two_any(t_pool, med) if len(t_pool) >= 2 else (None, None))
        level = nearest_one_any(l_pool, med) if l_pool else None
        meta["miss_t"] = station is None or back is None
        meta["miss_l"] = level is None
        return station, back, level, meta

    # R2: 无互通名 → 现行全局逻辑
    if target is None:
        meta["miss_t"] = meta["miss_l"] = True
        return None, None, None, meta
    station, back = nearest_two(traverse, target, prefix)
    level = nearest_one(levels, target, prefix)
    meta["path"] = "global"
    if station is None or back is None:
        meta["miss_t"] = True
    if level is None:
        meta["miss_l"] = True
    return station, back, level, meta


def copy_cell(src, dst):
    dst.value = src.value
    if src.has_style:
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)
        dst.alignment = copy(src.alignment)


def md5_file(path: Path) -> str:
    h = hashlib.md5()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def paint_row_yellow(ws, row, max_col):
    for c in range(1, max_col + 1):
        ws.cell(row, c).fill = YELLOW


def write_missing_list(path: Path, rows: list[dict], title: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "缺测站"
    headers = ["序号", "台账B原文", "清洗后B", "所在幅", "目标桩号", "缺导线", "缺水准", "无桩号", "备注"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    for i, r in enumerate(rows, 1):
        ws.cell(i + 1, 1, r["idx"])
        ws.cell(i + 1, 2, r["raw"])
        ws.cell(i + 1, 3, r["b"])
        ws.cell(i + 1, 4, r["side"])
        ws.cell(i + 1, 5, r.get("target"))
        ws.cell(i + 1, 6, "是" if r["miss_t"] else "")
        ws.cell(i + 1, 7, "是" if r["miss_l"] else "")
        ws.cell(i + 1, 8, "是" if r["no_stake"] else "")
        ws.cell(i + 1, 9, r.get("note", ""))
    ws.cell(len(rows) + 3, 1, title)
    wb.save(path)
    wb.close()


def init_sheet_from_template(tws, title: str, max_col: int):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in (1, 2):
        for c in range(1, max_col + 1):
            copy_cell(tws.cell(r, c), ws.cell(r, c))
    for c in range(1, max_col + 1):
        letter = get_column_letter(c)
        if tws.column_dimensions[letter].width:
            ws.column_dimensions[letter].width = tws.column_dimensions[letter].width
    for r in (1, 2):
        if tws.row_dimensions[r].height:
            ws.row_dimensions[r].height = tws.row_dimensions[r].height
    for merged in list(tws.merged_cells.ranges):
        if merged.min_col <= max_col and merged.max_row <= 2:
            try:
                ws.merge_cells(str(merged))
            except Exception:
                pass
    return wb, ws


def load_ledger_b(path: Path, sheet_name: str):
    lwb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lws = lwb[sheet_name]
    rows = []
    for row in lws.iter_rows(min_row=4, max_col=2, values_only=True):
        b = row[1]
        if b is None or str(b).strip() == "":
            continue
        rows.append(str(b).strip())
    lwb.close()
    return rows


def build_one(cfg, traverse, levels, twb, corridor_medians, out_dir: Path, project: str):
    """cfg keys documented in main()."""
    ledger_rows = load_ledger_b(cfg["ledger"], cfg["ledger_sheet"])
    print(f"[{cfg['tag']}] 台账有效行: {len(ledger_rows)}")

    tws = twb[cfg["template_sheet"]]
    wb, ws = init_sheet_from_template(tws, cfg["template_sheet"], cfg["max_col"])

    report = {
        "rows": 0,
        "missing_traverse": 0,
        "missing_level": 0,
        "far_traverse": 0,
        "far_level": 0,
        "zab_adjusted": 0,
        "side_inferred": [],
        "side_unknown": [],
        "no_stake": 0,
        "yellow_rows": 0,
        "max_near_m": MAX_NEAR_M,
        "samples": [],
        "notes": list(cfg.get("notes", [])),
        "corridor_named": 0,
        "corridor_stake_ok": 0,
        "corridor_stake_median_fill": 0,
        "corridor_no_stake_fill": 0,
        "corridor_no_stake_fail": 0,
        "path_global": 0,
        "cross_wrong": 0,  # B 互通 ≠ 所选测站互通
        "cross_checked": 0,
    }
    missing_rows = []

    for idx, raw_b in enumerate(ledger_rows, start=1):
        excel_row = idx + 2
        b_clean = clean_b(raw_b, cfg["eng_prefix"])
        side, side_note = detect_side(raw_b)
        stakes = parse_stakes(raw_b)

        if side_note.startswith("【推断】"):
            report["side_inferred"].append((idx, raw_b, side, side_note))
        elif side_note.startswith("【无法判定】"):
            report["side_unknown"].append((idx, raw_b, side_note))

        prefix = stakes[0][0] if stakes else ""
        nums = [s[1] for s in stakes]

        if len(nums) >= 2:
            target = (nums[0] + nums[1]) / 2.0
        elif len(nums) == 1:
            target = float(nums[0])
        else:
            target = None
            report["no_stake"] += 1

        station, back, level, meta = match_stations(
            raw_b, prefix, target, traverse, levels, corridor_medians
        )
        miss_t = meta["miss_t"]
        miss_l = meta["miss_l"]

        if meta["corridor"]:
            report["corridor_named"] += 1
        if meta["path"] in ("corridor_stake", "corridor_stake_median") and not miss_t and not miss_l:
            report["corridor_stake_ok"] += 1
        if meta["path"] == "corridor_stake_median" and not miss_t and not miss_l:
            report["corridor_stake_median_fill"] = report.get("corridor_stake_median_fill", 0) + 1
        if meta["path"] == "corridor_no_stake":
            if meta["no_stake_corridor_fill"] and not miss_t and not miss_l:
                report["corridor_no_stake_fill"] += 1
            else:
                report["corridor_no_stake_fail"] += 1
        if meta["path"] == "global":
            report["path_global"] += 1

        if miss_t:
            report["missing_traverse"] += 1
            if target is not None:
                report["far_traverse"] += 1
        if miss_l:
            report["missing_level"] += 1
            if target is not None:
                report["far_level"] += 1

        # 跨互通粗检：B 有互通名且测站位置描述属另一互通
        b_cor = meta["corridor"]
        if b_cor and station and not miss_t:
            report["cross_checked"] += 1
            s_cor = point_corridor(station)
            if s_cor and s_cor != b_cor:
                report["cross_wrong"] += 1

        z_val, aa_val, ab_val, zab_adj, zab_note = compute_zab(
            nums[:2] if nums else [], cfg["pick_ab"]
        )
        if zab_adj:
            report["zab_adjusted"] += 1

        def put(col, val):
            cell = ws.cell(excel_row, col, val)
            src = tws.cell(4, col)
            if src.font:
                cell.font = copy(src.font)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            return cell

        # A, B, C common
        put(1, idx)
        put(2, b_clean)
        put(3, side)
        put(4, cfg["D"])
        put(5, cfg["E"])
        put(6, cfg["F"])
        # G left empty for all three

        if cfg["tag"] == "JA-001":
            # N-S traverse 14-19; T,U level 20-21; Z,AA,AB 26-28; AF=32; AG=33; AM=39
            if station and back:
                put(14, station["name"])
                put(15, round(station["x"], 3))
                put(16, round(station["y"], 3))
                put(17, back["name"])
                put(18, round(back["x"], 3))
                put(19, round(back["y"], 3))
            if level:
                put(20, level["name"])
                put(21, round(level["z"], 3))
            if z_val is not None:
                put(26, z_val)
                put(27, aa_val)
                put(28, ab_val)
            put(32, f'=CONCATENATE("(",A{excel_row},")",B{excel_row},C{excel_row},E{excel_row})')
            if prefix:
                put(33, prefix)
            put(39, f'=CONCATENATE("(",A{excel_row},")",B{excel_row},C{excel_row},E{excel_row},"抽检")')

        elif cfg["tag"] == "JA-002":
            # I-R empty (no ledger attrs); S-X traverse 19-24; Y,Z level 25-26
            # AA,AB,AC 27-29; AD,AE,AF 30-32 constants; AG=33; AH=34; AI=35
            if station and back:
                put(19, station["name"])
                put(20, round(station["x"], 3))
                put(21, round(station["y"], 3))
                put(22, back["name"])
                put(23, round(back["x"], 3))
                put(24, round(back["y"], 3))
            if level:
                put(25, level["name"])
                put(26, round(level["z"], 3))
            if z_val is not None:
                put(27, z_val)
                put(28, aa_val)
                put(29, ab_val)
            put(30, cfg["ad"])
            put(31, cfg["ae"])
            put(32, cfg["af"])
            put(33, f'=CONCATENATE("(",A{excel_row},")",B{excel_row},C{excel_row},E{excel_row})')
            if prefix:
                put(34, prefix)
            put(35, f'=CONCATENATE("(",A{excel_row},")",B{excel_row},C{excel_row},E{excel_row},"抽检")')

        elif cfg["tag"] == "JA-003":
            # I empty (砼数量 no source); J-O traverse 10-15; P,Q level 16-17
            # R,S,T 18-20; U-W empty; X=24; Y=25; Z=26
            if station and back:
                put(10, station["name"])
                put(11, round(station["x"], 3))
                put(12, round(station["y"], 3))
                put(13, back["name"])
                put(14, round(back["x"], 3))
                put(15, round(back["y"], 3))
            if level:
                put(16, level["name"])
                put(17, round(level["z"], 3))
            if z_val is not None:
                put(18, z_val)
                put(19, aa_val)
                put(20, ab_val)
            put(24, f'=CONCATENATE("(",A{excel_row},")",B{excel_row},C{excel_row},E{excel_row})')
            if prefix:
                put(25, prefix)
            put(26, f'=CONCATENATE("(",A{excel_row},")",B{excel_row},C{excel_row},E{excel_row},"抽检")')

        if miss_t or miss_l:
            paint_row_yellow(ws, excel_row, cfg["max_col"])
            report["yellow_rows"] += 1
            note_bits = [x for x in [side_note, zab_note, meta.get("path"), meta.get("corridor")] if x]
            if meta.get("path") == "corridor_no_stake" and target is None:
                note_bits.append(f"corridor_median={meta.get('target_used')}")
            missing_rows.append(
                {
                    "idx": idx,
                    "raw": raw_b,
                    "b": b_clean,
                    "side": side,
                    "target": meta.get("target_used", target),
                    "miss_t": miss_t,
                    "miss_l": miss_l,
                    "no_stake": target is None,
                    "note": "; ".join(str(x) for x in note_bits),
                }
            )

        if idx in (1, 2, 3) or idx == len(ledger_rows):
            sample = {
                "idx": idx,
                "B": b_clean,
                "C": side,
                "prefix": prefix,
                "ZAB": (z_val, aa_val, ab_val),
                "station": station["name"] if station else None,
                "back": back["name"] if back else None,
                "level": level["name"] if level else None,
                "yellow": miss_t or miss_l,
                "side_note": side_note,
                "corridor": meta.get("corridor"),
                "path": meta.get("path"),
            }
            report["samples"].append(sample)

        report["rows"] += 1
        if idx % 500 == 0:
            print(f"  [{cfg['tag']}] wrote {idx}...")

    out_xlsx = out_dir / cfg["out_name"]
    out_rep = out_dir / cfg["rep_name"]
    out_miss = out_dir / cfg["miss_name"]
    out_dir.mkdir(parents=True, exist_ok=True)
    if out_xlsx.exists():
        out_xlsx.unlink()
    wb.save(out_xlsx)
    wb.close()

    write_missing_list(out_miss, missing_rows, f"{cfg['tag']} 缺测站清单（黄行）")

    digest = md5_file(out_xlsx)
    size = out_xlsx.stat().st_size
    lines = [
        f"输出: {out_xlsx}",
        f"文件名约定: {cfg['out_name']}",
        f"项目编号: {project}",
        f"写入行数: {report['rows']}",
        f"文件大小: {size}",
        f"MD5: {digest}",
        f"邻近上限(m): {report['max_near_m']}",
        f"缺导线点(不足2个邻近或超距/无桩号): {report['missing_traverse']}",
        f"缺水准点(无邻近或超距/无桩号): {report['missing_level']}",
        f"有桩号但导线超距/不足: {report['far_traverse']}",
        f"有桩号但水准超距/不足: {report['far_level']}",
        f"无桩号行: {report['no_stake']}",
        f"黄行(缺测站整行FFFF00): {report['yellow_rows']}",
        f"Z/AA/AB(或对应列)曾调整: {report['zab_adjusted']}",
        f"左右幅推断: {len(report['side_inferred'])}",
        f"左右幅无法判定: {len(report['side_unknown'])}",
        f"B含互通名行: {report['corridor_named']}",
        f"走廊有桩号配成: {report['corridor_stake_ok']}",
        f"走廊有桩号经中位数兜底: {report['corridor_stake_median_fill']}",
        f"走廊无桩号救回(黄灭): {report['corridor_no_stake_fill']}",
        f"走廊无桩号仍黄: {report['corridor_no_stake_fail']}",
        f"无互通名走全局: {report['path_global']}",
        f"跨互通粗检(已配且B有名): checked={report['cross_checked']} wrong={report['cross_wrong']}",
        f"缺测站清单: {out_miss}",
        "",
    ]
    if report["notes"]:
        lines.append("=== 说明备注 ===")
        lines.extend(report["notes"])
        lines.append("")
    lines.append("=== 抽检样本 ===")
    for s in report["samples"]:
        lines.append(str(s))
    lines.append("")
    lines.append("=== 左右无法判定样例(最多30) ===")
    for item in report["side_unknown"][:30]:
        lines.append(str(item))
    lines.append("")
    lines.append("=== 左右推断样例(最多30) ===")
    for item in report["side_inferred"][:30]:
        lines.append(str(item))
    out_rep.write_text("\n".join(lines), encoding="utf-8")

    summary = {
        "tag": cfg["tag"],
        "out": str(out_xlsx),
        "rep": str(out_rep),
        "miss": str(out_miss),
        "rows": report["rows"],
        "missing_traverse": report["missing_traverse"],
        "missing_level": report["missing_level"],
        "no_stake": report["no_stake"],
        "yellow_rows": report["yellow_rows"],
        "side_inferred": len(report["side_inferred"]),
        "side_unknown": len(report["side_unknown"]),
        "zab_adjusted": report["zab_adjusted"],
        "corridor_named": report["corridor_named"],
        "corridor_no_stake_fill": report["corridor_no_stake_fill"],
        "corridor_no_stake_fail": report["corridor_no_stake_fail"],
        "corridor_stake_ok": report["corridor_stake_ok"],
        "corridor_stake_median_fill": report["corridor_stake_median_fill"],
        "cross_wrong": report["cross_wrong"],
        "cross_checked": report["cross_checked"],
        "far_traverse": report["far_traverse"],
        "far_level": report["far_level"],
        "md5": digest,
        "size": size,
    }
    print(f"DONE {cfg['tag']}: rows={summary['rows']} yellow={summary['yellow_rows']} md5={digest}")
    return summary


class InputError(Exception):
    """数据目录或参数不完整，调用方应原样告诉操作的人。"""


ITEM_SPECS = {
    "JA-001": {
        "tag": "JA-001",
        "ledger_name": "资料台账-JA-001.xlsx",
        "instruction_name": "资料参数执行说明-JA-001.md",
        "ledger_sheet": "交安工程-JA-001",
        "template_sheet": "标志牌-JA-001",
        "stem": "标志牌-JA-001",
        "eng_prefix": "标志工程",
        "D": "交安工程",
        "E": "交通标志",
        "F": "交通标志",
        "pick_ab": pick_ab_ja001,
        "max_col": 39,
        "notes": [
            "G–M / V–Y / AC–AE / AH–AL 留空。",
            "缺 2 个导线点或缺水准点：对应测站格留空，整行黄填 FFFF00。",
            "互通走廊只认全称：白岭互通、石坳互通、港口互通、修水西枢纽、铜鼓枢纽。",
            "无桩号且有互通全称：用走廊主线 K 中位数作目标，池内取最近 2 个导线点和 1 个水准点，不套 2km。",
            "有桩号且有互通全称：池内同前缀优先，再池内任意前缀，仍不超过 2000m；池内不足则改用走廊主线 K 中位数，距离相对中位数仍不超过 2000m。",
            "无互通全称：全局同前缀优先，再任意前缀，不超过 2000m。不启用 3km 兜底。",
        ],
    },
    "JA-002": {
        "tag": "JA-002",
        "ledger_name": "资料台账-JA-002.xlsx",
        "instruction_name": "资料参数执行说明-JA-002.md",
        "ledger_sheet": "交安工程-JA-002",
        "template_sheet": "标线-JA-002",
        "stem": "标线-JA-002",
        "eng_prefix": "标线工程",
        "D": "交安工程",
        "E": "路面标线",
        "F": "安全设施",
        "ad": 2.2,
        "ae": 0.1,
        "af": 0.01,
        "pick_ab": pick_ab_ja002,
        "max_col": 35,
        "notes": [
            "说明第 7 条要求 I 至 R 列按源表数据填写。当前台账没有这些属性列，故 I–R 留空，不补造。",
            "AC 分档：小于等于 50 取 10；大于 50 且小于等于 70 取 20；大于 70 且小于等于 200 取 40；大于 200 取 100。单桩号 AC 为 0。",
            "AD/AE/AF 默认每行 2.2 / 0.1 / 0.01；执行说明里能解析出这三个数时改用说明。",
            "缺测站整行黄填 FFFF00。",
            "互通锁池与 2km 规则同 JA-001。不启用 3km 兜底。",
        ],
    },
    "JA-003": {
        "tag": "JA-003",
        "ledger_name": "资料台账-JA-003.xlsx",
        "instruction_name": "资料参数执行说明-JA-003.md",
        "ledger_sheet": "交安工程-JA-003",
        "template_sheet": "波形护栏-JA-003",
        "stem": "波形护栏-JA-003",
        "eng_prefix": "波形护栏工程",
        "D": "交安工程",
        "E": "波形护栏",
        "F": "安全设施",
        "pick_ab": pick_ab_ja001,
        "max_col": 26,
        "notes": [
            "执行说明里的表名「标线-JA-003」是笔误。样板工作表和输出文件名使用「波形护栏-JA-003」。",
            "说明第 10 条末尾「AA 和 AB 列」是笔误，按同条前文填写 R、S；只有一个桩号时 R=S，T=0。",
            "I 列砼数量在台账中没有来源，留空。U–W 留空。",
            "缺测站整行黄填 FFFF00。",
            "互通锁池与 2km 规则同 JA-001。不启用 3km 兜底。",
        ],
    },
}


def parse_def(text: str):
    match = DEF_RE.search(text or "")
    if not match:
        return None
    return tuple(part.strip() for part in match.groups())


def parse_ja002_constants(text: str):
    match = CONST_RE.search(text or "")
    if not match:
        return None
    return tuple(float(part) for part in match.groups())


def read_instruction(path: Path) -> str | None:
    if not path.is_file():
        return None
    return path.read_text(encoding="utf-8-sig").replace("\r\n", "\n").replace("\r", "\n")


def prepare_cfg(spec: dict, project: str, data_dir: Path) -> dict:
    cfg = dict(spec)
    stem = spec["stem"]
    cfg["ledger"] = data_dir / spec["ledger_name"]
    cfg["out_name"] = f"{stem}-{project}.xlsx"
    cfg["rep_name"] = f"{stem}-{project}-生成报告.txt"
    cfg["miss_name"] = f"{stem}-{project}-缺测站清单.xlsx"
    instruction_path = data_dir / spec["instruction_name"]
    text = read_instruction(instruction_path)
    notes = [
        "邻近上限 2000m。互通走廊只认全称。不启用 3km 兜底。",
    ]
    if text:
        notes.append(f"已读取执行说明: {instruction_path.name}")
        notes.append("工作表名、输出文件名和测站匹配使用内置规则。")
        parsed = parse_def(text)
        builtin = (spec["D"], spec["E"], spec["F"])
        if parsed:
            cfg["D"], cfg["E"], cfg["F"] = parsed
            if parsed == builtin:
                notes.append("D/E/F 与执行说明一致。")
            else:
                notes.append(
                    "D/E/F 改用执行说明："
                    + " / ".join(parsed)
                    + "。"
                )
        else:
            notes.append("执行说明中未解析出 D/E/F，使用内置值。")
        if spec["tag"] == "JA-002":
            constants = parse_ja002_constants(text)
            builtin_constants = (spec["ad"], spec["ae"], spec["af"])
            if constants:
                cfg["ad"], cfg["ae"], cfg["af"] = constants
                if constants == builtin_constants:
                    notes.append("AD/AE/AF 与执行说明一致。")
                else:
                    notes.append(
                        "AD/AE/AF 改用执行说明："
                        + " / ".join(str(item) for item in constants)
                        + "。"
                    )
            else:
                notes.append("执行说明中未解析出 AD/AE/AF，使用内置 2.2 / 0.1 / 0.01。")
    else:
        notes.append("数据目录没有执行说明，D/E/F 与常量使用内置规则。")
    notes.extend(spec["notes"])
    cfg["notes"] = notes
    return cfg


def resolve_template(data_dir: Path, explicit: Path | None = None) -> Path:
    """数据目录里的参数样板优先；没有则用仓库 share/参数样板.xlsx。"""
    if explicit is not None:
        path = explicit.expanduser().resolve()
        if not path.is_file():
            raise InputError(f"找不到参数样板: {path}")
        return path
    local = (data_dir / "参数样板.xlsx").resolve()
    if local.is_file():
        return local
    shared = CANONICAL_TEMPLATE.resolve()
    if shared.is_file():
        return shared
    raise InputError(
        "找不到参数样板.xlsx。可放在数据目录，或使用仓库 share/参数样板.xlsx。"
    )


def normalize_items(items: list[str]) -> list[str]:
    chosen = []
    seen = set()
    for raw in items:
        tag = raw.strip().upper()
        if tag not in ITEM_SPECS:
            known = "、".join(ITEM_SPECS)
            raise InputError(f"未知分项编号: {raw}。可选 {known}。")
        if tag not in seen:
            seen.add(tag)
            chosen.append(tag)
    if not chosen:
        raise InputError("至少指定一个分项编号。")
    return chosen


def normalize_project(project: str) -> str:
    text = (project or "").strip()
    if not PROJECT_RE.fullmatch(text):
        raise InputError("项目编号只允许字母、数字、下划线和短横线，长度 1 到 32。")
    return text


def _require_sheet(workbook, sheet_name: str, path: Path):
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    names = "、".join(workbook.sheetnames)
    raise InputError(f"{path.name} 中没有工作表「{sheet_name}」。现有工作表：{names}。")


def run_items(items: list[str], project: str, data_dir: Path, template: Path | None = None) -> list[dict]:
    tags = normalize_items(items)
    project = normalize_project(project)
    data_dir = data_dir.expanduser().resolve()
    if not data_dir.is_dir():
        raise InputError(f"数据目录不存在: {data_dir}")

    template_path = resolve_template(data_dir, template)
    survey_path = data_dir / "导线点及水准点.xlsx"
    missing = []
    if not survey_path.is_file():
        missing.append(survey_path.name)
    for tag in tags:
        ledger_path = data_dir / ITEM_SPECS[tag]["ledger_name"]
        if not ledger_path.is_file():
            missing.append(ledger_path.name)
    if missing:
        raise InputError("数据目录缺少必要文件：\n" + "\n".join(missing))

    print(f"数据目录: {data_dir}")
    print(f"参数样板: {template_path}")
    swb = openpyxl.load_workbook(survey_path, read_only=True, data_only=True)
    try:
        traverse = load_points(_require_sheet(swb, "导线点", survey_path), need_xy=True)
        levels = load_points(_require_sheet(swb, "水准点", survey_path), need_xy=False)
    finally:
        swb.close()
    print(f"导线点可用: {len(traverse)}，水准点可用: {len(levels)}")

    corridor_medians = {}
    for cname in CORRIDOR_CANON:
        med = corridor_k_median(traverse, cname)
        corridor_medians[cname] = med
        t_n = len(corridor_pool(traverse, cname))
        l_n = len(corridor_pool(levels, cname))
        print(f"  走廊 {cname}: 导线池={t_n} 水准池={l_n} K中位={med}")

    twb = openpyxl.load_workbook(template_path, data_only=False)
    out_dir = data_dir / "output"
    summaries = []
    try:
        for tag in tags:
            if ITEM_SPECS[tag]["template_sheet"] not in twb.sheetnames:
                _require_sheet(twb, ITEM_SPECS[tag]["template_sheet"], template_path)
            cfg = prepare_cfg(ITEM_SPECS[tag], project, data_dir)
            summaries.append(build_one(cfg, traverse, levels, twb, corridor_medians, out_dir, project))
    finally:
        twb.close()

    lines = [
        "交安资料参数生成汇总",
        f"项目编号: {project}",
        f"参数样板: {template_path}",
        f"邻近上限(m): {MAX_NEAR_M}",
        "互通走廊只认全称：白岭互通、石坳互通、港口互通、修水西枢纽、铜鼓枢纽。",
        "有桩号时不超过 2000m；无桩号且有互通全称时用走廊主线 K 中位数，不套 2000m。",
        "不启用 3km 兜底。",
        "",
    ]
    for item in summaries:
        lines.append(
            f"{item['tag']}: rows={item['rows']} yellow={item['yellow_rows']} "
            f"no_stake={item['no_stake']} far_t={item['far_traverse']} "
            f"miss_t={item['missing_traverse']} miss_l={item['missing_level']} "
            f"corridor_no_stake_fill={item['corridor_no_stake_fill']} "
            f"corridor_stake_median_fill={item['corridor_stake_median_fill']} "
            f"cross_wrong={item['cross_wrong']}/{item['cross_checked']} "
            f"size={item['size']} md5={item['md5']}"
        )
        lines.append(f"  xlsx={item['out']}")
        lines.append(f"  rep={item['rep']}")
        lines.append(f"  miss={item['miss']}")
    summary_path = out_dir / "生成汇总.txt"
    summary_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(summary_path.read_text(encoding="utf-8"))
    return summaries
